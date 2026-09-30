r"""Generate a Project Health Check .docx for the project from state.json.

Usage: python render_health_docx.py [--project ID] [--out PATH]
Default output: <state folder>\reports\<project name> - Project Health Check - YYYY-MM-DD.docx
"""

from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
import guard
import pm_state
import project

STATUS_COLOURS = {
    "current": RGBColor(0x27, 0x50, 0x0A),
    "stale": RGBColor(0x85, 0x4F, 0x0B),
    "missing": RGBColor(0xA3, 0x2D, 0x2D),
    "not_required": RGBColor(0x5F, 0x5E, 0x5A),
}


dd = pm_state.display_dates


def shade(cell, hex_fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


REPORT = {"byline": "", "currency": "$", "area_unit": "sqm", "cost_label": "Budget"}


def money(v):
    return REPORT["currency"] + format(round(v), ",") if v is not None else "-"


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        for r in c.paragraphs[0].runs:
            r.bold = True
        shade(c, "F1EFE8")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = dd(val) if val is not None else "-"
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    return t


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", help="registered project id (default: PM_PROJECT, or the only one)")
    ap.add_argument("--out")
    args = ap.parse_args()

    proj = project.select(args.project)
    cfg = pm_state.load_config()
    REPORT.update(cfg.get("report", {}))
    milestone = (cfg.get("programme") or {}).get("milestone", "Milestone")
    unit = REPORT["area_unit"]

    state = pm_state.load_state()
    store = pm_state.project_state(state)
    if not store:
        raise SystemExit(f"{proj.name} has not been scanned yet - run scan_state.py first")

    name = store["name"]
    today_iso = date.today().isoformat()
    today = dd(today_iso)
    dh = store.get("doc_health", {})
    issue = (store.get("drawing_stage") or {}).get("latest_issue") or {}
    prog = store.get("program", {})
    cost = store.get("cost", {})
    governing = prog.get("confirmed_date") or prog.get("governing_date")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    doc.add_heading(f"Project Health Check — {name}", level=0)
    byline = f"{REPORT['byline']} · " if REPORT["byline"] else ""
    p = doc.add_paragraph(
        f"{byline}Generated {today} by the PM agent from a read-only scan of the project "
        f"folder. Scan time: {dd(store.get('last_scanned', '')[:16])}.")
    p.runs[0].font.size = Pt(9.5)

    doc.add_heading("Summary", level=1)
    band = dh.get("band", "?").upper()
    s = doc.add_paragraph()
    r = s.add_run(f"Documentation health: {dh.get('score', '?')}% ({band})")
    r.bold = True
    r.font.color.rgb = STATUS_COLOURS.get(
        {"GREEN": "current", "AMBER": "stale", "RED": "missing"}.get(band, "not_required"),
        STATUS_COLOURS["not_required"])
    doc.add_paragraph(
        f"Drawing stage: {dh.get('stage', '?')} — {issue.get('set_file', 'no set found')} "
        f"(rev {issue.get('revision') or '?'}, issued {dd(issue.get('issue_date', '?'))}). "
        f"Disciplines in set: {', '.join(issue.get('disciplines_detected') or ['not verified'])}.")
    waiting = store.get("waiting_on", [])
    doc.add_paragraph(
        f"{len(waiting)} open item{'s' if len(waiting) != 1 else ''} · "
        f"{REPORT['cost_label']} {money(cost.get('total'))} ({money(cost.get('per_area'))}/{unit}, "
        f"{cost.get('area') or '?'} {unit}) · {milestone.lower()} "
        f"{dd(governing) if governing else 'unconfirmed'}"
        f"{' (conflicting dates on file)' if prog.get('conflict') and not prog.get('confirmed_date') else ''}.")

    doc.add_heading("Intake documentation", level=1)
    rows = []
    for docname, d in dh.get("docs", {}).items():
        rows.append((docname.replace("_", " ").capitalize(), d["status"], d.get("detail") or ""))
    t = add_table(doc, ["Document", "Status", "Detail"], rows)
    fills = {"current": "EAF3DE", "stale": "FAEEDA", "missing": "FCEBEB", "not_required": "F1EFE8"}
    for i, (_, status, _) in enumerate(rows, start=1):
        shade(t.rows[i].cells[1], fills.get(status, "FFFFFF"))

    doc.add_heading("Open items", level=1)
    if waiting:
        add_table(doc, ["Type", "Item", "With", "Raised"],
                  [(w.get("type", "").replace("_", " "), w.get("text"),
                    w.get("needed_from"), w.get("raised")) for w in waiting])
    else:
        doc.add_paragraph("None - all clear.")

    resolved = store.get("resolved", {})
    if resolved:
        doc.add_heading("Resolved this period", level=1)
        add_table(doc, ["Item", "Resolution"],
                  [(rid.removeprefix(f"{proj.id}-"),
                    f"{v.get('answer', '')} — {v.get('resolution', '')}".strip(" —"))
                   for rid, v in resolved.items()])

    notes = store.get("standing_notes", [])
    if notes:
        doc.add_heading("Standing instructions", level=1)
        for n in notes:
            doc.add_paragraph(n, style="List Bullet")

    doc.add_heading("Programme and cost position", level=1)
    add_table(doc, ["Measure", "Value"], [
        (f"{milestone} (governing)", governing or "unconfirmed"),
        (f"{milestone} dates on file", ", ".join(prog.get("milestone_dates", [])) or "-"),
        ("Latest programme file", (prog.get("latest") or {}).get("name", "-")),
        (REPORT["cost_label"], money(cost.get("total"))),
        ("Rate", f"{money(cost.get('per_area'))}/{unit} over {cost.get('area') or '?'} {unit}"),
        ("Contract sum", money(cost.get("contract_sum")) if cost.get("contract_sum") else "pre-award"),
        ("Commitments on file", f"{cost.get('purchase_orders', 0)} POs · {cost.get('invoices', 0)} invoices · {cost.get('quotes', 0)} quotes"),
    ])

    footer = doc.add_paragraph()
    fr = footer.add_run(
        "Statuses: current = dated on/after the latest drawing issue · stale = predates it · "
        "missing = required at this stage, not found. Source: the PM's state.json.")
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = RGBColor(0x6E, 0x6A, 0x60)
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT

    out = guard.check(args.out or pm_state.REPORTS_DIR / f"{name} - Project Health Check - {today_iso}.docx",
                      "write")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    pm_state.log_activity("health_report", str(out))
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
