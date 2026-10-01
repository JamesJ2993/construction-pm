r"""Project scanner: reads the project folder (READ-ONLY) and updates state.json.

Usage:
  python scan_state.py [--project ID]           # scan, skipping if nothing changed
  python scan_state.py --dry-run                # print proposed state, write nothing
  python scan_state.py --deep                   # also PyMuPDF-scan the latest set for disciplines
  python scan_state.py --force                  # rescan even if the folder fingerprint is unchanged

Every folder name, filename pattern and label it looks for comes from the project's config.json -
see reference/setup.md. Never writes anywhere except the project's state folder.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import guard
import pm_state
import project
from pm_state import make_waiting, now_iso

RE_REV_SUFFIX = re.compile(r"[\s_-](P|T|C)(\d+)(?=[\s._)]|$)", re.I)
RE_REV_ALPHA = re.compile(r"rev[\s._-]*([A-Z])(?![a-z])", re.I)
RE_DATE6 = re.compile(r"(?<!\d)(2\d{5})(?!\d)")
RE_DATE8 = re.compile(r"(?<!\d)(20\d{6})(?!\d)")
RE_DMY = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{2}(?:\d{2})?)")

DEFAULT_SKIP_DIRS = ["CADs", "CAD", "LINKEDs", "_SS", "SS", "pages-to-jpg"]
DEFAULT_PREFIXES = {"A": "architectural", "S": "structural", "M": "mechanical",
                    "E": "electrical", "H": "hydraulic", "F": "fire"}
DISCIPLINE_KEYWORDS = {"structural": "structural", "mechanical": "mechanical",
                       "electrical": "electrical", "hydraulic": "hydraulic",
                       "fire": "fire", "sprinkler": "fire"}
PLACEHOLDERS = {"thumbs.db", "desktop.ini", ".ds_store"}


def ignored(name: str, cfg: dict) -> bool:
    if name in cfg.get("ignore_files", []):
        return True
    return any(name.startswith(p) for p in cfg.get("ignore_prefixes", ["~$"]))


def list_files(folder: Path, cfg: dict, exts=None, depth: int = 3) -> list[Path]:
    """Bounded recursive listing, skipping CAD/link/system clutter. The folder is guard-checked."""
    if not folder.is_dir():
        return []
    guard.check(folder, "read")
    out = []
    skip_dirs = set(cfg.get("skip_dirs", DEFAULT_SKIP_DIRS))

    def walk(d: Path, level: int):
        try:
            entries = list(d.iterdir())
        except OSError:
            return
        for e in entries:
            try:
                if e.is_dir():
                    if level < depth and e.name not in skip_dirs:
                        walk(e, level + 1)
                elif not ignored(e.name, cfg):
                    if exts is None or e.suffix.lower() in exts:
                        out.append(e)
            except OSError:
                continue

    walk(folder, 0)
    return out


def mtime_iso(p: Path) -> str:
    return datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds")


def parse_revision(name: str) -> str | None:
    m = RE_REV_SUFFIX.search(name)
    if m:
        return (m.group(1) + m.group(2)).upper()
    m = RE_REV_ALPHA.search(name)
    if m:
        return m.group(1).upper()
    return None


def parse_filename_date(name: str) -> str | None:
    for rx, fmt in ((RE_DATE8, "%Y%m%d"), (RE_DATE6, "%y%m%d")):
        m = rx.search(name)
        if m:
            try:
                return datetime.strptime(m.group(1), fmt).date().isoformat()
            except ValueError:
                pass
    return None


def parse_dmy(text: str) -> list[str]:
    """DD.MM.YY or DD.MM.YYYY dates in a filename, as ISO."""
    out = []
    for d, m, y in RE_DMY.findall(text):
        y = ("20" + y) if len(y) == 2 else y
        try:
            out.append(datetime(int(y), int(m), int(d)).date().isoformat())
        except ValueError:
            continue
    return out


def _exts(values) -> set[str]:
    return {e.lower() for e in values}


# ---------------------------------------------------------------- drawings

def scan_drawings(root: Path, cfg: dict) -> dict:
    dcfg = cfg["drawings"]
    drawings = root / dcfg.get("dir", "Drawings")
    skip = set(dcfg.get("skip_folders", []))

    result = {"current": None, "latest_issue": None, "stages_found": {}}
    best = None
    for st in dcfg["stages"]:
        folder_name = st["folder"]
        stage_dir = drawings / folder_name
        if not stage_dir.is_dir():
            continue
        pdfs = list_files(stage_dir, cfg, exts={".pdf"})
        result["stages_found"][folder_name] = len(pdfs)
        if not pdfs or not st.get("is_issue", True) or folder_name in skip:
            continue
        best = (st, pdfs)
    if best is None:
        return result

    st, pdfs = best
    pinned = dcfg.get("pinned_set_pattern")
    pool = pdfs
    if pinned:
        hits = [p for p in pdfs if pinned.lower() in p.name.lower()]
        if hits:
            pool = hits
    latest_pdf = max(pool, key=lambda p: p.stat().st_mtime)
    result["current"] = st["stage"]
    result["latest_issue"] = {
        "stage_folder": st["folder"],
        "set_file": latest_pdf.name,
        "set_path": str(latest_pdf),
        "issue_date": parse_filename_date(latest_pdf.name) or mtime_iso(latest_pdf)[:10],
        "modified": mtime_iso(latest_pdf),
        "revision": parse_revision(latest_pdf.name),
        "pinned": bool(pinned and pool is not pdfs),
        "disciplines_detected": None,
        "disciplines_missing": None,
    }
    return result


def deep_scan_disciplines(set_path: str, cfg: dict) -> list[str] | None:
    """Classify sheets in a combined set by sheet-number prefix, then title-block keywords."""
    try:
        import fitz
    except ImportError:
        return None
    guard.check(set_path, "read")
    prefixes = cfg["drawings"].get("discipline_prefixes", DEFAULT_PREFIXES)
    sheet_rx = re.compile(r"\b([" + "".join(prefixes) + r"])[\s-]?\d{3,5}\b")
    found = set()
    try:
        doc = fitz.open(set_path)
        for page in doc:
            rect = page.rect
            clip = fitz.Rect(rect.width * 0.6, rect.height * 0.6, rect.width, rect.height)
            text = page.get_text(clip=clip).upper()
            m = sheet_rx.search(text)
            if m:
                found.add(prefixes[m.group(1)])
                continue
            low = text.lower()
            for kw, disc in DISCIPLINE_KEYWORDS.items():
                if kw in low:
                    found.add(disc)
                    break
            else:
                found.add("architectural")
        doc.close()
    except Exception:
        return None
    return sorted(found) if found else None


# ---------------------------------------------------------------- documents

def find_registers(root: Path, cfg: dict) -> list[dict]:
    rcfg = cfg["registers"]
    pattern = re.compile(rcfg.get("pattern", "clarific|register"), re.I)
    exclude = re.compile(rcfg["exclude"], re.I) if rcfg.get("exclude") else None
    seen, regs = set(), []
    for rel in rcfg.get("dirs", []):
        for f in list_files(root / rel, cfg, exts=_exts(rcfg.get("exts", [".xlsx", ".xlsm"])),
                            depth=rcfg.get("depth", 1)):
            if pattern.search(f.name) and not (exclude and exclude.search(f.name)):
                if str(f) not in seen:
                    seen.add(str(f))
                    regs.append({"path": str(f), "name": f.name,
                                 "modified": mtime_iso(f), "size": f.stat().st_size,
                                 "revision": parse_revision(f.name)})
    regs.sort(key=lambda r: r["modified"])
    return regs


def find_file(root: Path, cfg: dict, searches: list[dict]) -> dict | None:
    """Try each search in order; the newest match of the first search with any hit wins."""
    for s in searches:
        pat = re.compile(s["pattern"], re.I)
        hits = [f for f in list_files(root / s["dir"], cfg, exts=_exts(s["exts"]),
                                      depth=s.get("depth", 1))
                if pat.search(f.name)]
        if hits:
            f = max(hits, key=lambda p: p.stat().st_mtime)
            return {"path": str(f), "name": f.name, "modified": mtime_iso(f)}
    return None


def folder_populated(root: Path, cfg: dict, names: list[str], exclude_latest: str | None = None) -> dict | None:
    """Newest file in the first named folder (or a folder whose name starts with it) that has any."""
    for name in names:
        d = root / name
        if not d.is_dir() and root.is_dir():
            for child in root.iterdir():
                if child.is_dir() and child.name.lower().startswith(name.lower()):
                    d = child
                    break
        if d.is_dir():
            files = list_files(d, cfg, depth=2)
            if files:
                latest = max(files, key=lambda p: p.stat().st_mtime)
                if exclude_latest and re.search(exclude_latest, latest.name, re.I):
                    return None
                return {"count": len(files), "latest": latest.name, "modified": mtime_iso(latest)}
    return None


def scan_documents(root: Path, cfg: dict) -> dict:
    arts = {}
    regs = find_registers(root, cfg)
    arts["registers"] = regs
    arts["register"] = {"found": True, **regs[-1]} if regs else {"found": False}
    for doc, spec in cfg.get("documents", {}).items():
        if "find" in spec:
            arts[doc] = find_file(root, cfg, spec["find"])
        elif "folder" in spec:
            arts[doc] = folder_populated(root, cfg, spec["folder"], spec.get("exclude_latest"))
    return arts


# ---------------------------------------------------------------- programme and cost

def scan_programme(root: Path, cfg: dict) -> dict:
    pcfg = cfg.get("programme") or {}
    result = {"found": False, "latest": None, "milestone_dates": [], "conflict": False}
    if not pcfg:
        return result
    files = list_files(root / pcfg["dir"], cfg, exts=_exts(pcfg.get("exts", [".pdf", ".mpp"])), depth=1)
    if not files:
        return result
    result["found"] = True
    latest = max(files, key=lambda p: p.stat().st_mtime)
    result["latest"] = {"name": latest.name, "modified": mtime_iso(latest)}
    dates = sorted({d for f in files for d in parse_dmy(f.name)})
    result["milestone_dates"] = dates
    result["conflict"] = len(dates) > 1
    governing = parse_dmy(latest.name)
    if governing:
        result["governing_date"] = governing[-1]
    return result


def read_budget(root: Path, cfg: dict) -> dict:
    bcfg = cfg.get("budget") or {}
    labels = bcfg.get("labels", {})
    result = {"found": False, "file": None, "per_area": None, **{k: None for k in labels}}
    if not bcfg:
        return result
    fin = root / bcfg["dir"]
    for name, sub in bcfg.get("count_dirs", {}).items():
        result[name] = len(list_files(fin / sub, cfg, depth=1))
    pattern = re.compile(bcfg.get("file_pattern", "budget|cost"), re.I)
    budgets = [f for f in list_files(fin, cfg, exts={".xlsx"}, depth=0) if pattern.search(f.name)]
    if not budgets:
        return result
    f = max(budgets, key=lambda p: p.stat().st_mtime)
    result["found"] = True
    result["file"] = {"name": f.name, "modified": mtime_iso(f)}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(guard.check(f, "read"), data_only=True, read_only=True)
        for ws in wb.worksheets[:3]:
            for row in ws.iter_rows(max_row=200, max_col=15):
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    label = cell.value.strip().lower()
                    if "per sqm" in label or "per m2" in label:
                        continue
                    for key, pats in labels.items():
                        if result[key] is None and any(label.startswith(p.lower()) for p in pats):
                            for c2 in row[cell.column:]:
                                if isinstance(c2.value, (int, float)):
                                    result[key] = round(float(c2.value), 2)
                                    break
        wb.close()
        if result.get("total") and result.get("area"):
            result["per_area"] = round(result["total"] / result["area"], 2)
    except Exception as e:
        result["read_error"] = str(e)[:200]
    return result


# ---------------------------------------------------------------- evidence register

def scan_evidence_register(root: Path, cfg: dict) -> dict | None:
    """Reconcile a requirements matrix against its numbered evidence folders.

    Rows carry an item number and a Status (Open / Closed). Evidence for item N sits in a sibling
    folder whose name starts with N. An Open item with evidence on file is a close candidate; a
    Closed item whose folder is empty is unsupported.
    """
    ecfg = cfg.get("evidence_register")
    if not ecfg:
        return None
    base = root / ecfg["dir"]
    if not base.is_dir():
        return None
    fpat = re.compile(ecfg.get("file_pattern", "matrix"), re.I)
    matrices = [f for f in list_files(base, cfg, exts={".xlsx"}, depth=2) if fpat.search(f.name)]
    if not matrices:
        return None
    matrix = max(matrices, key=lambda p: p.stat().st_mtime)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(guard.check(matrix, "read"), data_only=True, read_only=True)
        spat = re.compile(ecfg.get("sheet_pattern", "."), re.I)
        sheet = next((s for s in wb.sheetnames if spat.search(s)), None)
        if sheet is None:
            wb.close()
            return None
        rows = list(wb[sheet].iter_rows(max_row=ecfg.get("max_rows", 200), max_col=20, values_only=True))
        wb.close()
    except Exception:
        return None
    hdr_i = next((i for i, r in enumerate(rows)
                  if r and any(v and "status" in str(v).lower() for v in r)), None)
    if hdr_i is None:
        return None
    hdr = [str(v).strip().lower() if v else "" for v in rows[hdr_i]]
    si = next(i for i, h in enumerate(hdr) if "status" in h)
    item_col, req_col = ecfg.get("item_col", 0), ecfg.get("requirement_col", 2)
    items = {}
    for r in rows[hdr_i + 1:]:
        if r and r[item_col] is not None and str(r[item_col]).strip().isdigit():
            items[int(str(r[item_col]).strip())] = {
                "requirement": str(r[req_col] or "")[:80],
                "status": str(r[si] or "").strip() or "blank"}

    placeholders = PLACEHOLDERS | {p.lower() for p in ecfg.get("placeholders", [])}
    evidence = {}
    for d in matrix.parent.iterdir():
        m = re.match(r"(\d+)", d.name)
        if d.is_dir() and m:
            files = [f for f in list_files(d, cfg, depth=2) if f.name.lower() not in placeholders]
            evidence[int(m.group(1))] = [f.name for f in files]

    close_candidates = [
        {"item": n, "requirement": it["requirement"], "files": evidence.get(n, [])}
        for n, it in items.items()
        if it["status"].lower() == "open" and evidence.get(n)]
    unsupported_closed = [
        {"item": n, "requirement": it["requirement"]}
        for n, it in items.items()
        if it["status"].lower() == "closed" and n in evidence and not evidence.get(n)]
    counts = {"total": len(items),
              "open": sum(1 for i in items.values() if i["status"].lower() == "open"),
              "closed": sum(1 for i in items.values() if i["status"].lower() == "closed")}
    return {"label": ecfg.get("label", "Requirements"), "matrix": matrix.name,
            "matrix_modified": mtime_iso(matrix), "counts": counts,
            "close_candidates": close_candidates, "unsupported_closed": unsupported_closed}


# ---------------------------------------------------------------- health, waiting-on, pipeline

def evaluate_doc_health(entry: dict, cfg: dict) -> dict:
    first_issue = next((s["stage"] for s in cfg["drawings"]["stages"] if s.get("is_issue", True)), None)
    stage = entry["drawing_stage"]["current"] or first_issue
    reqs = cfg["required_documents"]
    fallback = next(iter(reqs.values()), [])
    ignore = set(cfg.get("ignore_docs", []))
    stale_ok = set(cfg.get("stale_ok", []))
    date_sensitive = set(cfg.get("date_sensitive_docs", ["register", "checklist", "sow", "tender_analysis"]))
    required = [d for d in reqs.get(stage, fallback) if d not in ignore]
    weights = cfg.get("doc_weights", {})
    expected = cfg["drawings"].get("expected_disciplines", [])
    issue = entry["drawing_stage"].get("latest_issue") or {}
    issue_mod = issue.get("modified", "")
    arts = entry["artifacts"]
    arts_alias = {"program": entry["program"].get("latest") if entry["program"]["found"] else None,
                  "budget": entry["cost"]["file"] if entry["cost"]["found"] else None}

    def status_for(doc: str) -> tuple[str, str]:
        if doc == "drawing_set":
            if not issue:
                return "missing", "no issued drawing set found"
            discs = issue.get("disciplines_detected")
            if discs is not None and expected and len(discs) < len(expected):
                missing = [d for d in expected if d not in discs]
                return "stale", "set incomplete: missing " + ", ".join(missing)
            return "current", issue.get("set_file", "")
        art = arts_alias[doc] if doc in arts_alias else arts.get(doc)
        if not art or (isinstance(art, dict) and not art.get("found", True)):
            return "missing", "not found"
        mod = art.get("modified", "") if isinstance(art, dict) else ""
        if doc in date_sensitive and mod and issue_mod and mod < issue_mod:
            if doc in stale_ok:
                return "current", f"dated {mod[:10]} - staleness waived by {pm_state.owner()}"
            return "stale", f"dated {mod[:10]}, before latest issue {issue_mod[:10]}"
        return "current", mod[:10]

    docs, score, total = {}, 0.0, 0.0
    for doc in required:
        st, detail = status_for(doc)
        docs[doc] = {"status": st, "detail": detail}
        w = weights.get(doc, 1)
        total += w
        score += w * (1.0 if st == "current" else 0.5 if st == "stale" else 0.0)
    pct = round(100 * score / total) if total else 0
    return {"stage": stage, "docs": docs, "score": pct,
            "band": "green" if pct >= 75 else "amber" if pct >= 50 else "red"}


def build_waiting(key: str, entry: dict, cfg: dict) -> list[dict]:
    owner = pm_state.owner()
    team = f"{owner} / project team"
    milestone = (cfg.get("programme") or {}).get("milestone", "Milestone")
    waiting = []
    issue = entry["drawing_stage"].get("latest_issue") or {}
    discs = issue.get("disciplines_detected")
    if discs is not None and "architectural" in discs and len(discs) == 1:
        waiting.append(make_waiting(f"{key}-disciplines", "missing_input",
            "Drawing set is architectural only - services disciplines not issued",
            "design consultants"))
    regs = entry["artifacts"].get("registers", [])
    if len(regs) >= 2 and regs[-1]["size"] < 0.2 * regs[-2]["size"]:
        waiting.append(make_waiting(f"{key}-register-size", "question",
            f"{regs[-1]['name']} is {regs[-1]['size']//1024} KB vs previous {regs[-2]['size']//1024} KB"
            " - embedded snapshots may be lost; verify before use", owner))
    for doc, info in entry.get("doc_health", {}).get("docs", {}).items():
        if info["status"] == "missing":
            waiting.append(make_waiting(f"{key}-{doc}", "missing_input",
                f"{doc.replace('_', ' ')}: {info['detail']}", team))
        elif info["status"] == "stale":
            waiting.append(make_waiting(f"{key}-{doc}", "missing_input",
                f"{doc.replace('_', ' ')} is stale - {info['detail']}", team))
    sow = entry["artifacts"].get("sow")
    if sow and regs and sow["modified"] < regs[-1]["modified"]:
        waiting.append(make_waiting(f"{key}-sow-superseded", "question",
            f"Scope of Works ({sow['modified'][:10]}) predates register {regs[-1].get('revision') or ''} "
            f"({regs[-1]['modified'][:10]}) - written against superseded documentation", owner))
    stages = [s["stage"] for s in cfg["drawings"]["stages"]]
    tender = cfg["drawings"].get("tender_stage", "tender")
    current = entry["drawing_stage"].get("current")
    if sow and current in stages and tender in stages and stages.index(current) < stages.index(tender):
        waiting.append(make_waiting(f"{key}-tender-on-prelim", "question",
            "A Scope of Works exists but the latest issued drawings are still "
            f"{entry['drawing_stage']['current']} - confirm the tender basis", owner))
    ev = entry.get("evidence")
    if ev:
        label = ev["label"]
        for c in ev.get("close_candidates", []):
            waiting.append(make_waiting(f"{key}-evidence-item-{c['item']}", "question",
                f"{label} item {c['item']} ({c['requirement'][:50]}) has evidence on file "
                f"({', '.join(c['files'][:2])}) but is still marked Open - close it?", owner))
        for c in ev.get("unsupported_closed", []):
            waiting.append(make_waiting(f"{key}-evidence-unsupported-{c['item']}", "question",
                f"{label} item {c['item']} ({c['requirement'][:50]}) is marked Closed but its "
                "evidence folder is empty - refile the evidence or reopen the item", owner))
    prog = entry.get("program", {})
    if prog.get("conflict"):
        waiting.append(make_waiting(f"{key}-milestone-date", "question",
            f"Conflicting {milestone.lower()} dates in the programme folder: "
            + ", ".join(prog["milestone_dates"])
            + (f" - latest file says {prog.get('governing_date')}" if prog.get("governing_date") else ""),
            owner))
    return waiting


def derive_pipeline(entry: dict) -> dict:
    dh = entry.get("doc_health", {}).get("docs", {})
    issue = entry["drawing_stage"].get("latest_issue")

    def doc_status(d):
        return dh.get(d, {}).get("status")

    pipeline = {}
    if not issue:
        pipeline["drawing_intake"] = {"status": "blocked", "detail": "no issued set found"}
    elif doc_status("drawing_set") != "current":
        pipeline["drawing_intake"] = {"status": "blocked",
                                      "detail": dh.get("drawing_set", {}).get("detail", "set incomplete")}
    else:
        pipeline["drawing_intake"] = {"status": "signed_off", "detail": "set current (auto-verified)"}
    regs = entry["artifacts"].get("registers", [])
    if not regs:
        pipeline["register_review"] = {"status": "pending", "detail": "no register - initial build needed"}
    elif doc_status("register") == "stale":
        pipeline["register_review"] = {"status": "pending", "detail": "register predates latest issue - re-review due"}
    else:
        pipeline["register_review"] = {"status": "in_review", "detail": "register current - PM review then sign-off"}
    pipeline["register_signoff"] = {"status": "pending", "detail": "not yet presented"}
    pipeline["doc_checklist"] = {"status": "pending" if doc_status("checklist") != "current" else "in_review",
                                 "detail": dh.get("checklist", {}).get("detail")}
    pipeline["sow_check"] = {"status": "blocked", "detail": "gated on register sign-off"}
    pipeline["programme_review"] = {"status": "pending", "detail": None}
    pipeline["cost_review"] = {"status": "pending", "detail": None}
    for step in pipeline.values():
        step["updated"] = now_iso()
    return pipeline


def watched_dirs(cfg: dict) -> list[str]:
    """Every folder the scan reads, for the change fingerprint."""
    dirs = [cfg["drawings"].get("dir", "Drawings"), *cfg["registers"].get("dirs", [])]
    for spec in cfg.get("documents", {}).values():
        dirs += [s["dir"] for s in spec.get("find", [])] + list(spec.get("folder", []))
    for section in ("programme", "budget", "evidence_register"):
        if cfg.get(section):
            dirs.append(cfg[section]["dir"])
    return sorted(set(dirs))


def fingerprint(root: Path, cfg: dict) -> str:
    h = hashlib.sha1()
    for sub in watched_dirs(cfg):
        for f in sorted(list_files(root / sub, cfg, depth=2), key=lambda p: str(p)):
            try:
                h.update(f"{f.name}|{f.stat().st_mtime_ns}".encode("utf-8", "replace"))
            except OSError:
                continue
    return h.hexdigest()


def scan_project(proj: project.Project, cfg: dict, deep: bool = False, prev: dict | None = None) -> dict:
    root = guard.check(proj.root, "read")
    entry = {"name": proj.name, "path": str(proj.root), "aliases": cfg.get("aliases", [proj.name])}

    entry["drawing_stage"] = scan_drawings(root, cfg)
    issue = entry["drawing_stage"].get("latest_issue")
    if deep and issue:
        discs = deep_scan_disciplines(issue["set_path"], cfg)
        if discs is not None:
            issue["disciplines_detected"] = discs
            issue["disciplines_missing"] = [d for d in cfg["drawings"].get("expected_disciplines", [])
                                            if d not in discs]
    elif issue and prev:
        # A deep result stays valid while the set file is unchanged.
        old = (prev.get("drawing_stage") or {}).get("latest_issue") or {}
        if (old.get("set_path"), old.get("modified")) == (issue["set_path"], issue["modified"]):
            for k in ("disciplines_detected", "disciplines_missing"):
                issue[k] = old.get(k)

    entry["artifacts"] = scan_documents(root, cfg)
    entry["program"] = scan_programme(root, cfg)
    entry["cost"] = read_budget(root, cfg)
    entry["evidence"] = scan_evidence_register(root, cfg)
    entry["doc_health"] = evaluate_doc_health(entry, cfg)
    entry["waiting_on"] = build_waiting(proj.id, entry, cfg)
    entry["pipeline"] = derive_pipeline(entry)
    entry["last_scanned"] = now_iso()
    entry["scan_fingerprint"] = fingerprint(root, cfg)
    return entry


def merge(prev: dict, fresh: dict) -> dict:
    """Carry forward what the owner decided; let the scan refresh everything it measures."""
    fresh["pipeline"] = prev.get("pipeline", fresh["pipeline"])
    fresh["resolved"] = prev.get("resolved", {})
    if prev.get("standing_notes"):
        fresh["standing_notes"] = prev["standing_notes"]
    if prev.get("program", {}).get("confirmed_date"):
        fresh["program"]["confirmed_date"] = prev["program"]["confirmed_date"]
    merged = {w["id"]: w for w in fresh["waiting_on"]}
    for w in prev.get("waiting_on", []):
        if w["id"] in merged:
            if w.get("type") == "external" or w.get("expected"):
                merged[w["id"]] = {**merged[w["id"]], **w}
        elif w.get("type") in ("approval", "external", "statutory"):
            merged[w["id"]] = w
    fresh["waiting_on"] = [w for w in merged.values() if w["id"] not in fresh["resolved"]]
    return fresh


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", help="registered project id (default: PM_PROJECT, or the only one)")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--deep", action="store_true", help="PyMuPDF discipline scan of latest set")
    ap.add_argument("--force", action="store_true", help="rescan even if fingerprint unchanged")
    args = ap.parse_args()

    proj = project.select(args.project)
    cfg = pm_state.load_config()
    state = pm_state.load_state()
    prev = state["projects"].get(proj.id)

    if prev and not (args.force or args.deep or args.dry_run):
        if prev.get("scan_fingerprint") == fingerprint(proj.root, cfg):
            print("scanned=0 skipped=1 (unchanged)")
            return 0

    fresh = scan_project(proj, cfg, deep=args.deep, prev=prev)
    if prev:
        fresh = merge(prev, fresh)

    if args.dry_run:
        print(json.dumps(fresh, indent=2, default=str))
        return 0
    state["projects"][proj.id] = fresh
    pm_state.save_state(state)
    pm_state.log_activity("scan", f"{proj.name} scanned", deep=args.deep)
    print(f"scanned=1 waiting_on={len(fresh['waiting_on'])} health={fresh['doc_health']['score']}%")
    return 0


if __name__ == "__main__":
    sys.exit(main())
