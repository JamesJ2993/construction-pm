#!/usr/bin/env python3
"""Build a client-ready .docx from a working markdown deliverable.

House rule: markdown is the working format; the thing that gets sent is a .docx
with a title page, on the firm's letterhead when it has one. This script does the
first two parts deterministically; ``apply-letterhead.py`` confirms the letterhead
afterwards. Never re-apply a letterhead to an issued document.

Usage:
    python md-to-docx.py <source.md> <target.docx> --title "..." [options]

    --title T      Document title on the title page (required).
    --subtitle S   Second line on the title page. Repeatable.
    --meta K=V     Title-page metadata row. Repeatable, order preserved.
    --status S     Status banner on the title page (e.g. "DRAFT - for review").
    --skip-h1      Drop the source's first H1 (it duplicates --title).
    --template P   Base .docx supplying styles, headers and footers - the firm's
                   letterhead (config.json -> report.letterhead). The document
                   inherits its typography rather than python-docx's stock base.
    --config P     A project config.json; its report.letterhead is the template.
                   With neither, the stock python-docx base is used.

Supports: ATX headings, pipe tables, bullet/numbered lists, blockquotes,
horizontal rules (page-break suppressed), **bold**, *italic* and `code` spans.
Deterministic: no network, no model, no judgement -- a format converter only.
"""

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor

INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`)")


def add_runs(para, text):
    """Render inline bold / italic / code spans into a paragraph."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            para.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = para.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            para.add_run(piece[1:-1]).italic = True
        else:
            para.add_run(piece)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_divider(line):
    cells = split_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != "")


def title_page(doc, title, subtitles, meta, status):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(26)
    for sub in subtitles:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(sub)
        run.font.size = Pt(14)
    if status:
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(status)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
    if meta:
        doc.add_paragraph()
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key, value in meta:
            cells = table.add_row().cells
            r = cells[0].paragraphs[0].add_run(key)
            r.bold = True
            add_runs(cells[1].paragraphs[0], value)
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


def build(src, dst, title, subtitles, meta, status, skip_h1, template=None):
    lines = Path(src).read_text(encoding="utf-8").splitlines()
    if template:
        doc = Document(str(template))
        # The template supplies styles, headers, footers and margins only --
        # drop any body content it carries so the build starts clean.
        body = doc.element.body
        for child in list(body):
            if not child.tag.endswith("}sectPr"):
                body.remove(child)
    else:
        doc = Document()
    title_page(doc, title, subtitles, meta, status)

    i, dropped_h1 = 0, not skip_h1
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Table: a pipe row followed by a divider row.
        if (stripped.startswith("|") and i + 1 < len(lines)
                and is_divider(lines[i + 1])):
            header = split_row(stripped)
            i += 2
            body = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                body.append(split_row(lines[i]))
                i += 1
            width = max([len(header)] + [len(r) for r in body])
            table = doc.add_table(rows=1, cols=width)
            table.style = "Table Grid"
            for c, text in enumerate(header):
                para = table.rows[0].cells[c].paragraphs[0]
                add_runs(para, text)
                for run in para.runs:
                    run.bold = True
            for row in body:
                cells = table.add_row().cells
                for c, text in enumerate(row[:width]):
                    add_runs(cells[c].paragraphs[0], text)
            doc.add_paragraph()
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level, text = len(heading.group(1)), heading.group(2)
            if level == 1 and not dropped_h1:
                dropped_h1 = True
                i += 1
                continue
            add_runs(doc.add_heading(level=min(level, 4)), text)
            i += 1
            continue

        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            i += 1
            continue

        if stripped.startswith(">"):
            add_runs(doc.add_paragraph(style="Intense Quote"),
                     stripped.lstrip("> ").strip())
            i += 1
            continue

        bullet = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if bullet:
            add_runs(doc.add_paragraph(style="List Bullet"), bullet.group(1))
            i += 1
            continue

        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if numbered:
            add_runs(doc.add_paragraph(style="List Number"), numbered.group(1))
            i += 1
            continue

        # Paragraph: gather until a blank line or a block starter.
        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^\s*(#{1,6}\s|[-*+]\s|\d+[.)]\s|\||>|-{3,}$)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        if buf:
            add_runs(doc.add_paragraph(), " ".join(buf))
        else:
            i += 1

    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    doc.save(dst)
    return dst


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("target")
    ap.add_argument("--title", required=True)
    ap.add_argument("--subtitle", action="append", default=[])
    ap.add_argument("--meta", action="append", default=[])
    ap.add_argument("--status", default="")
    ap.add_argument("--skip-h1", action="store_true")
    ap.add_argument("--template", default=None,
                    help="Base .docx for styles/headers - the firm's letterhead")
    ap.add_argument("--config", default=None,
                    help="Project config.json; report.letterhead is used as --template")
    ap.add_argument("--no-template", action="store_true",
                    help="Build on python-docx's stock base even if a letterhead is configured")
    args = ap.parse_args(argv)

    template = None
    if not args.no_template:
        if args.template:
            template = Path(args.template)
        elif args.config:
            import json
            cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
            lh = (cfg.get("report") or {}).get("letterhead")
            template = Path(lh) if lh else None
        if template is not None and not template.is_file():
            sys.exit(f"letterhead not found: {template}")

    meta = []
    for item in args.meta:
        if "=" not in item:
            sys.exit(f"--meta expects KEY=VALUE, got: {item}")
        key, value = item.split("=", 1)
        meta.append((key.strip(), value.strip()))

    out = build(args.source, args.target, args.title, args.subtitle, meta,
                args.status, args.skip_h1, template=template)
    base = template.name if template else "stock python-docx base"
    print(f"{Path(out).name}: built (base: {base})")


if __name__ == "__main__":
    main()
