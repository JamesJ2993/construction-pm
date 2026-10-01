#!/usr/bin/env python3
"""check-provenance.py - did this deliverable come from its controlled structure, or was it invented?

Presentation gates test appearance: is it on the letterhead, does it open on a contents page. A
document invented from nothing passes both perfectly. This gate asks the other question - does the
artefact descend from the controlled template or report structure, and did that structure survive
being filled.

Two sources of structure, in this order:

  1. The firm's template library - config.json -> templates.dir. The deliverable is matched to a
     template by name, and the template's own Heading 1/2 texts (for .docx) or sheet names and
     upper-case section labels (for .xlsx) must survive in the deliverable. Markers are read FROM
     the template, never from memory of what a section is called elsewhere.
  2. The plugin's report structures - skills/deliverable-build/structures/*.md. Each declares its
     required_sections, written with terminology placeholders ({{payment_application}}, {{change}})
     that are rendered in the project's own terms first, so a US pay application assessment is
     checked against US headings and an Australian progress claim assessment against Australian ones.

USAGE
    python check-provenance.py --check <deliverable> --config <state>/config.json [--structure NAME]
    python check-provenance.py --check <deliverable> --exempt "inbound contractor document"
    python check-provenance.py --list --config <state>/config.json

    --structure names the structure or template explicitly (the commission says which). Without it
    the deliverable is matched by file name.

EXIT CODES
    0  check-pass
    1  check-fail (no matching template or structure, or its structure did not survive)
    2  usage or environment error - INCLUDING a configured templates.dir that does not exist. A
       missing library is an error, never a silent fallback to something else.

WHAT IT DOES NOT DO
    Judge content, arithmetic or quality. Independent review and the owner's sign-off stay where
    they are.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
STRUCTURES = HERE.parent / "structures"
sys.path.insert(0, str(HERE.parent.parent / "knowledge-base" / "scripts"))

EXEMPT_NAMES = [
    (r"(?i)\bREAD ?ME\b", "working note, not a deliverable"),
    (r"(?i)\bTEMPLATE\b", "is itself a template"),
]
FM = re.compile(r"\A---\s*\n(.*?)\n---\s*\n", re.S)
PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z_]+)\s*\}\}")


# ----------------------------------------------------------------------------- terms
def load_terms(cfg: dict) -> dict:
    facts = cfg.get("project_facts") or {}
    terms: dict = {}
    country = facts.get("country")
    if country:
        try:
            import kbpaths
            terms = {k: v for k, v in kbpaths.load_json(country, "terminology.json").items()
                     if not k.startswith("_")}
        except Exception:  # noqa: BLE001 - terms are best effort; markers fall back to key names
            terms = {}
    terms.update({k: v for k, v in (cfg.get("terminology") or {}).items() if v})
    return terms


def render(text: str, terms: dict) -> str:
    def sub(m):
        key = m.group(1)
        val = terms.get(key.lower()) or key.lower().replace("_", " ")
        if key.isupper():
            return val.upper()
        if key[0].isupper():
            return val[:1].upper() + val[1:]
        return val
    return PLACEHOLDER.sub(sub, text)


def norm(s: str) -> str:
    """Lower-case, parentheticals dropped, whitespace squashed - for comparing headings."""
    s = re.sub(r"\([^)]*\)", " ", s or "")
    s = s.replace("&", " and ")
    return " ".join(re.sub(r"[^\w\s/]", " ", s.lower()).split())


# ----------------------------------------------------------------------------- reading
def parse_frontmatter(text: str) -> dict:
    m = FM.match(text)
    meta: dict = {}
    if not m:
        return meta
    key = None
    for line in m.group(1).splitlines():
        if line.strip().startswith("#") or not line.strip():
            continue
        item = re.match(r"^\s+-\s+(.*)$", line)
        if item and key:
            meta.setdefault(key, [])
            if isinstance(meta[key], list):
                meta[key].append(item.group(1).strip().strip("\"'"))
            continue
        if ":" not in line:
            continue
        k, _, v = line.partition(":")
        key, v = k.strip(), v.strip()
        if v.startswith("[") and v.endswith("]"):
            meta[key] = [x.strip() for x in v[1:-1].split(",") if x.strip()]
        elif v:
            meta[key] = v.strip("\"'")
        else:
            meta[key] = []
    return meta


def docx_headings(path: Path, levels=("Title", "Heading 1", "Heading 2", "Heading 3")) -> list[str]:
    import docx
    d = docx.Document(str(path))
    return [p.text.strip() for p in d.paragraphs
            if p.style is not None and p.style.name in levels and p.text.strip()]


def docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "ignore")
    return re.sub(r"<[^>]+>", " ", xml)


def workbook(path: Path):
    try:
        import openpyxl
    except ImportError:
        sys.stderr.write("check-provenance: openpyxl is required for .xlsx checks\n")
        raise SystemExit(2)
    return openpyxl.load_workbook(path, read_only=True, data_only=False)


def deliverable_name(path: Path) -> str:
    """Strip job number, claim/revision suffixes and dates, leaving the artefact name."""
    stem = path.stem
    stem = re.sub(r"^[A-Z]{1,6}-?\d{1,6}\s*[-—]\s*", "", stem)
    stem = re.sub(r"\s*[-—]\s*(PC|VO|VAR|CO|PA)-?\d+\b.*$", "", stem, flags=re.I)
    stem = re.sub(r"\s*[-—]\s*\d{4}-\d{2}-\d{2}\b.*$", "", stem)
    stem = re.sub(r"\s*[-—]\s*Rev\s*[A-Z0-9]+\b.*$", "", stem, flags=re.I)
    return stem.strip()


# ----------------------------------------------------------------------------- sources
def library_templates(library: Path) -> dict:
    out = {}
    for p in sorted(library.rglob("*")):
        if p.suffix.lower() not in (".docx", ".xlsx", ".dotx", ".xltx") or p.name.startswith("~$"):
            continue
        name = re.sub(r"\s*[-—]?\s*TEMPLATE$", "", deliverable_name(p), flags=re.I).strip()
        out.setdefault(name, p)
    return out


def structures(terms: dict) -> dict:
    out = {}
    for p in sorted(STRUCTURES.glob("*.md")):
        meta = parse_frontmatter(p.read_text(encoding="utf-8"))
        if not meta.get("deliverable"):
            continue
        meta["_path"] = p
        names = [meta["deliverable"]] + list(meta.get("aliases") or [])
        for n in names:
            out.setdefault(render(n, terms), meta)
    return out


def match(artefact: str, candidates: dict):
    a = norm(artefact)
    for name, val in candidates.items():
        if norm(name) == a:
            return name, val
    for name, val in candidates.items():
        n = norm(name)
        if n and (n in a or a in n):
            return name, val
    return None


# ----------------------------------------------------------------------------- checks
def check_against_template(path: Path, tpl: Path) -> list[str]:
    failures = []
    kind = path.suffix.lower()
    if kind == ".docx" and tpl.suffix.lower() in (".docx", ".dotx"):
        have = {norm(h) for h in docx_headings(path)} | {norm(docx_text(path))}
        blob = " ".join(have)
        for marker in docx_headings(tpl, levels=("Heading 1", "Heading 2")):
            m = norm(re.sub(r"\[[^\]]*\]", " ", marker))
            if m and m not in blob:
                failures.append(f"missing template section: {marker!r}")
    elif kind == ".xlsx" and tpl.suffix.lower() in (".xlsx", ".xltx"):
        twb, dwb = workbook(tpl), workbook(path)
        for sheet in twb.sheetnames:
            if re.search(r"(?i)copy me|instructions|guidance", sheet):
                continue
            if sheet not in dwb.sheetnames:
                failures.append(f"missing template sheet: {sheet!r}")
                continue
            labels = {str(r[0].value).strip() for r in twb[sheet].iter_rows(min_col=1, max_col=1, max_row=120)
                      if isinstance(r[0].value, str) and r[0].value.strip().isupper() and len(r[0].value.strip()) > 3}
            have = {str(r[0].value).strip().upper() for r in dwb[sheet].iter_rows(min_col=1, max_col=1)
                    if isinstance(r[0].value, str)}
            for label in sorted(labels):
                if label.upper() not in have:
                    failures.append(f"sheet {sheet!r} has lost template section {label!r}")
        twb.close()
        dwb.close()
    else:
        failures.append(f"deliverable is {kind} but the matched template is {tpl.suffix}")
    return failures


def check_against_structure(path: Path, meta: dict, terms: dict) -> list[str]:
    failures = []
    kind = path.suffix.lower()
    if kind == ".docx":
        heads = [norm(h) for h in docx_headings(path)]
        for section in meta.get("required_sections") or []:
            want = norm(render(section, terms))
            if want and not any(want in h for h in heads):
                failures.append(f"missing required section heading: {render(section, terms)!r}")
    elif kind == ".xlsx":
        sheets = meta.get("required_sheets") or []
        if sheets:
            wb = workbook(path)
            for s in sheets:
                if not any(norm(render(s, terms)) in norm(n) for n in wb.sheetnames):
                    failures.append(f"missing required sheet: {render(s, terms)!r}")
            wb.close()
    return failures


def check(path: Path, cfg: dict, structure: str | None, exempt: str | None, verbose=True) -> int:
    if exempt:
        print(f"{path.name}: check-pass (exempt - {exempt})")
        return 0
    for pattern, why in EXEMPT_NAMES:
        if re.search(pattern, path.name):
            print(f"{path.name}: check-pass (exempt - {why})")
            return 0

    terms = load_terms(cfg)
    artefact = structure or deliverable_name(path)
    lib = (cfg.get("templates") or {}).get("dir")
    if lib:
        library = Path(lib)
        if not library.is_dir():
            sys.stderr.write(f"check-provenance: templates.dir is set but not found: {library}\n"
                             f"  This is an ERROR, not a pass - fix the config or the path.\n")
            raise SystemExit(2)
        found = match(artefact, library_templates(library))
        if found:
            name, tpl = found
            failures = check_against_template(path, tpl)
            return report(path, failures, f"template {tpl.name}")

    found = match(render(artefact, terms), structures(terms))
    if found is None:
        print(f"{path.name}: check-FAIL")
        print(f"    No template or report structure matches {artefact!r}.")
        print("    A deliverable with no controlled structure was invented, not filled. Build it from")
        print("    the named structure, or record why it is exempt with --exempt \"<reason>\".")
        return 1
    _name, meta = found
    return report(path, check_against_structure(path, meta, terms), f"structure {meta['_path'].name}")


def report(path: Path, failures: list[str], source: str) -> int:
    if failures:
        print(f"{path.name}: check-FAIL ({source})")
        print("    Traces to a controlled structure, but the structure did not survive:")
        for f in failures:
            print(f"      - {f}")
        return 1
    print(f"{path.name}: check-pass ({source})")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check", metavar="PATH")
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--config", metavar="CONFIG_JSON")
    ap.add_argument("--structure", metavar="NAME")
    ap.add_argument("--exempt", metavar="REASON")
    args = ap.parse_args()

    cfg = json.loads(Path(args.config).read_text(encoding="utf-8")) if args.config else {}
    if args.list:
        terms = load_terms(cfg)
        lib = (cfg.get("templates") or {}).get("dir")
        if lib and Path(lib).is_dir():
            print(f"template library: {lib}")
            for name, p in library_templates(Path(lib)).items():
                print(f"  {name:<50} {p.suffix}")
        print(f"report structures ({STRUCTURES}):")
        seen = set()
        for name, meta in structures(terms).items():
            if meta["_path"] in seen:
                continue
            seen.add(meta["_path"])
            secs = ", ".join(render(s, terms) for s in meta.get("required_sections") or [])
            print(f"  {render(meta['deliverable'], terms):<50} {secs}")
        return 0
    if args.check:
        p = Path(args.check)
        if not p.is_file():
            sys.stderr.write(f"check-provenance: not a file: {p}\n")
            return 2
        return check(p, cfg, args.structure, args.exempt)
    ap.print_help()
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
